# -*- coding: utf-8 -*-
"""
B / C 端成品导出器 · JSON 规则库 -> Excel 交付表
--------------------------------------------------
用法：
  python export_excel.py --preset gd                     # B 端 广东·珠三角 9 市（区域包）
  python export_excel.py --preset gd --scope province    # **广东省总表**（跨城归并 + 覆盖率标注）
  python export_excel.py --preset all                    # B 端 已收录城市（全库，标题由 PRESET_TITLES 推导）
  python export_excel.py --cities SZ,GZ,DG --out my.xlsx

  # C 端（2026-10-12 新增）：只收「个人 / 创业者 / 海外人才」可申报
  python export_excel.py --audience c --preset prd       # 珠三角 9 市个人补贴清单
  python export_excel.py --audience c --preset all       # 已收录城市个人补贴清单

产出：
  B 端 output/<标题>政策补贴总表_<版本>.xlsx
    Sheet1 政策总表   — 24 列，含自动筛选，可直接给 HR / 财税 / 培训机构用
    Sheet2 字段说明   — 每列含义与用法
    Sheet3 数据字典   — 枚举中英对照 + 溯源覆盖统计（含免责声明）
  C 端 output/C端/<标题>_<版本>.xlsx
    Sheet1 政策总表   — 20 列，**城市列前置 + 自动筛选**（多城靠筛选定位，不必线性翻）
    Sheet2 按城市速览 — 城市 / 条数 / 类别分布 / 编号区间（第一跳索引）
    Sheet3 字段说明
    Sheet4 数据字典与声明
源目录全程只读。

★ C 端为什么是 Excel 而不是 PDF（2026-10-12 用户实测反馈）：
  「单城适合 PDF；多城几百条＋每条都有条件/材料/入口，PDF 得线性翻，找某城市的政策要
    先记政策名再找城市，很麻烦。」→ 多城改出 Excel：城市列可筛选、可排序、可搜索。
  收录口径复用 `audience.is_person`（与 C 端 PDF 同一份判据，不另写）。
"""
import argparse
import collections
import datetime
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from audience import is_person            # C 端收录口径（与 C 端 PDF 同源，见 pipeline/audience.py）
import cross_city
import geo
import loader
import link_guard as LG
import mappers as M

# ★ C 端收录口径的「第三类主体」文案 = `enum_map.applicant_kind.overseas`（与 export_pdf 同源，2026-10-17 C 项）
_OS_KIND = M.t("applicant_kind", "overseas") or "海外人才"
_OS_SCOPE = "个人 / 创业者 / %s" % _OS_KIND

VERSION = "v20261017"                  # 成品文件名里的版本号：改了数据要自己换一个，
                                       # 否则两批成品同名、分不清哪份是新的。
                                       # 换号规则（自定即可）：重新对外交付就换一个日期号，
                                       # 只修错别字之类不换。★ 必须与 export_pdf.VERSION 同号。
SHEET_VERSION = "表结构 v1.2"          # v1.1 起：新增「政策原文链接」列（含可点击超链接）
                                       # v1.2 起：新增「区县配套对照」列（district_awards 字段转正）
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(color="FFFFFF", bold=True, size=10, name="微软雅黑")
BODY_FONT = Font(size=10, name="微软雅黑")
LINK_FONT = Font(size=9, name="微软雅黑", color="0563C1", underline="single")
ALT_FILL = PatternFill("solid", fgColor="F2F5FA")
THIN = Side(style="thin", color="D0D7E5")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

URL_COLS = {"办理入口", "政策原文链接"}
# L5（2026-10-09）：列名 → 链接角色。降级判据按角色区分（guide_url 还要判文种），
# 见 pipeline/link_guard.py；判据真值源是源库 collector/link_audit.json。
URL_ROLE = {"办理入口": "apply_url", "政策原文链接": "guide_url"}
AUDIT = LG.load_audit()
LINK_STRICT = False      # --link-strict：连栏目/聚合页、解读页、新闻页也降级
DOWNGRADED = []          # 本次出成品被降级的链接，收尾写台账
CENTER_COLS = {"序号", "政策编号", "省份", "城市", "适用对象",
               "补贴形式", "兑现方式", "办理渠道", "数据深度", "更新时间"}

# (列名, 列宽, 是否换行)
COLS = [
    ("序号", 6, False),
    ("政策编号", 13, False),
    ("省份", 7, False),
    ("城市", 7, False),
    ("政策类别", 14, False),
    ("细分类型", 16, False),
    ("政策名称", 34, True),
    ("标准名称", 26, True),
    ("适用对象", 14, False),
    ("补贴形式", 10, False),
    ("兑现方式", 16, False),
    ("办理渠道", 13, False),
    ("补贴标准", 52, True),
    ("区县配套对照", 46, True),
    ("申报窗口", 26, True),
    ("申报条件", 52, True),
    ("所需材料", 46, True),
    ("办理入口", 40, True),
    ("渠道说明", 34, True),
    ("政策依据", 44, True),
    ("政策原文链接", 44, True),
    ("有效期", 20, True),
    ("数据深度", 11, False),
    ("更新时间", 11, False),
]

FIELD_NOTE = [
    ("序号", "本表行号，非政策编号"),
    ("政策编号", "唯一编号，可直接引用给客户或用于对账"),
    ("省份 / 城市", "归属地，用于地域筛选"),
    ("政策类别", "9 大类之一，用于大类筛选"),
    ("细分类型", "类别下的业务细分，如「企业用工补贴」"),
    ("政策名称", "政策事项全称（以官方原文为准）"),
    ("标准名称", "标准化后的简称，便于去重与跨城对比"),
    ("适用对象", "个人 / 用人单位 / 创业者"),
    ("补贴形式", "现金补贴 / 资格通道 / 评审认定"),
    ("兑现方式", "核准兑现 / 免申即享 / 统筹名额 / 排序竞争 / 专家评审"),
    ("办理渠道", "线上办理 / 线下窗口 / 线上线下均可"),
    ("补贴标准", "金额与计算口径原文表述；「据实」「分档」类以官方审核为准"),
    ("区县配套对照", "同一事项在本市下辖各区县/镇街的差异化配套标准（首次认定 / 重新认定 / "
                    "迁入口径 / 依据）。市级统一标准见「补贴标准」列；留空表示该政策无区县差异化配套。"
                    "凡标「依据」的均可回原文核对，各区口径以当年申报指南为准"),
    ("申报窗口", "申报节奏（常年受理 / 按季度 / 按批次 / 集中申报）"),
    ("申报条件", "核心申报条件，含基础条件、择一分支（组内需同时满足、组间满足其一）与积分规则；"
                 "完整规则见官方文件"),
    ("所需材料", "材料清单，编号排列"),
    ("办理入口", "在线申报 / 办事指南入口链接（官方域名，可点击）。★ **留空 = 该政策无统一"
                "在线申报页面**（可能为线下窗口受理，或经电子邮箱 / 政务 APP 等渠道提交），"
                "并非数据缺失；具体办理路径见「办理渠道」与「渠道说明」"),
    ("渠道说明", "线上线下具体路径补充说明；留空表示该政策原文未细述办理路径"),
    ("政策依据", "文号与文件名，用于溯源核对"),
    ("政策原文链接", "政策原文 / 法规库页面链接（可点击）。优先取政策文件页；"
                    "少数条目原文与办事入口同页时会与「办理入口」相同"),
    ("有效期", "政策有效期或失效口径"),
    ("数据深度", "L0 仅目录 → L4 全字段，L4 为最高可用级"),
    ("更新时间", "该条政策最近核验日期"),
]

# ── C 端列集（2026-10-12）────────────────────────────────────────────────
# 与 B 端（24 列）的差异，全部围绕「个人读者自己找政策」：
#   · 加「适用范围」—— 多城清单里省级/国家级政策会在多座城市各出现一次，标注后才不会被
#     误会成重复条目（文案真值源同 C 端 PDF：`cross_city.scope_labels`）
#   · 去「细分类型 / 标准名称 / 兑现方式 / 区县配套对照 / 渠道说明」—— 机构侧字段
#   · **城市列前置**（第 4 列）+ 冻结窗格 + 自动筛选：点一下筛选箭头就是某座城市的清单
#     （这正是用户弃用多城 PDF 的原因：几百条 PDF 线性翻、找政策要先记名再找城）
C_COLS = [
    ("序号", 6, False),
    ("政策编号", 13, False),
    ("省份", 7, False),
    ("城市", 7, False),
    ("适用范围", 11, False),
    ("政策类别", 14, False),
    ("政策名称", 34, True),
    ("适用对象", 14, False),
    ("补贴形式", 10, False),
    ("补贴标准", 52, True),
    ("申报条件", 52, True),
    ("所需材料", 46, True),
    ("办理渠道", 13, False),
    ("办理入口", 40, True),
    ("政策原文链接", 44, True),
    ("政策依据", 44, True),
    ("申报窗口", 26, True),
    ("有效期", 20, True),
    ("数据深度", 11, False),
    ("更新时间", 11, False),
]

C_FIELD_NOTE = [
    ("序号", "本表行号，非政策编号"),
    ("政策编号", "唯一编号，可引用给客户或用于对账"),
    ("省份 / 城市", "归属地。★ **城市列可筛选** —— 只想看某座城市，点该列筛选箭头选城市即可"),
    ("适用范围", "该政策管到哪一级：「某市」= 该市本级政策；「某省全省」「全国通用」= 上级政策。"
                "★ **多城清单**用它区分层级（省级 / 国家级政策各市都会各自发布一次，本清单按城市"
                "分别列出、**不是重复条目**）；★ **单城清单**整本同城、标了是噪音，故该列留空"),
    ("政策类别", "9 大类之一，用于大类筛选"),
    ("政策名称", "政策事项全称（以官方原文为准）"),
    ("适用对象", f"{_OS_SCOPE} —— 本表只收录**个人可申报**的政策"),
    ("补贴形式", "现金补贴 / 资格通道 / 评审认定"),
    ("补贴标准", "金额与计算口径原文表述；「据实」「分档」类以官方审核为准"),
    ("申报条件", "核心申报条件，含基础条件、择一分支（组内需同时满足、组间满足其一）与积分规则；"
                 "完整规则见官方文件"),
    ("所需材料", "材料清单，编号排列"),
    ("办理渠道", "线上办理 / 线下窗口 / 线上线下均可"),
    ("办理入口", "在线申报 / 办事指南入口链接（官方域名，可点击）。★ **留空 = 该政策无统一"
                "在线申报页面**（可能为线下窗口受理，或经电子邮箱 / 政务 APP 等渠道提交），"
                "并非数据缺失；具体办理路径见「办理渠道」"),
    ("政策原文链接", "政策原文 / 法规库页面链接（可点击）。优先取政策文件页；"
                    "少数条目原文与办事入口同页时会与「办理入口」相同"),
    ("政策依据", "文号与文件名，用于溯源核对"),
    ("申报窗口", "申报节奏（常年受理 / 按季度 / 按批次 / 集中申报）"),
    ("有效期", "政策有效期或失效口径"),
    ("数据深度", "L0 仅目录 → L4 全字段（这条政策的信息完整度）；可筛选 L3/L4 只看信息最全的"),
    ("更新时间", "该条政策最近核验日期"),
]

C_CENTER_COLS = {"序号", "政策编号", "省份", "城市", "适用范围", "适用对象",
                 "补贴形式", "办理渠道", "数据深度", "更新时间"}

DICT_ROWS = [
    ("补贴形式", "money", "现金补贴"),
    ("补贴形式", "channel", "资格通道"),
    ("补贴形式", "review", "评审认定"),
    ("兑现方式", "approval", "核准兑现（符合即享）"),
    ("兑现方式", "auto", "免申即享"),
    ("兑现方式", "planned", "统筹名额"),
    ("兑现方式", "competition", "排序竞争"),
    ("兑现方式", "review", "专家评审"),
    ("办理渠道", "online", "线上办理"),
    ("办理渠道", "window", "线下窗口"),
    ("办理渠道", "mixed", "线上线下均可"),
    ("适用对象", "person", "个人"),
    ("适用对象", "unit", "用人单位"),
    ("适用对象", "entrepreneur", "创业者"),
    ("数据深度", "L0-L4", "政策信息的完整度分级，L4 最全"),
]


# ── 标题推导 + 范围校验（2026-10-10 修「优先级陷阱」）──────────────────
# 预设包的展示名。**不是**城市列表的真值源（那是 loader.PRESETS），只用于标题。
# ★ 2026-10-13：`all` 由「全国」改为「已收录城市」。
#   事故原型：全库成品名为「全国政策补贴总表」，但库里只有 **4 省 15 城** ——
#   B 端客户是机构（拿去报价、做方案），「全国」二字会造成**覆盖度误读**，
#   且这不是「数字会腐」那一类（名字里没写计数），而是**地理主张无据**。
#   判据：**产品名自称的地理范围必须有数据支撑**（见 `check_title_scope`）——
#   「全国」须覆盖 ≥ NATIONAL_MIN_PROVINCES 个省级行政区，否则拦红。
#   注意 `PRESET_TITLES` 只描述**已采集范围**，不描述行政区划本身。
PRESET_TITLES = {"gd": "广东（珠三角）", "yrd": "长三角", "gx": "广西", "all": "已收录城市"}


def derive_title(explicit, preset, cities, recs, scope="region"):
    """推导成品标题。

    **修的是这个 bug（2026-10-10）**：原写法
        title = args.title or f"{M.province(recs[0]['city'])}省" if recs else (args.title or "政策")
    因 `or` 优先级高于条件表达式，实际解析为 `(A or B) if C else D` ——
    漏 `--title` 时拿「排序后第一条记录」的省份当整表标题，
    于是 `--preset all`（全国 807 条）产出「广东省政策补贴总表」。
    **根因：标题取自排序副产物，而非声明的范围。**

    两条通道分开，避免区域包冒充省（珠三角 9 市 ≠ 广东省，两者会撞同名文件）：
      - `scope=province` → 一律「{省}省」
      - `scope=region`   → 用区域包名（PRESET_TITLES），**不从单省反推「X省」**
    """
    if explicit:
        return explicit
    if scope == "province":
        provs = sorted({M.province(c) for c in (cities or [])} - {""})
        return f"{provs[0]}省" if len(provs) == 1 else "政策"
    if preset in PRESET_TITLES:
        return PRESET_TITLES[preset]
    if not cities:
        # 无 preset 名 + 无城市列表 = 全库；**不得自称「全国」**（库里未必覆盖全国，
        # 见 `check_title_scope` 判据二）。用不含地理主张的说法。
        return "已收录城市"
    provs = sorted({M.province(c) for c in cities} - {""})
    return f"{provs[0]}省" if len(provs) == 1 else "政策"


NATIONAL_MIN_PROVINCES = 30   # 自称「全国」所需的最低省级行政区覆盖数（全国共 34 个）


def check_title_scope(title, recs):
    """防复发拦网：标题自称的范围必须容得下实际记录范围。

    判据一（2026-10-10）：标题写「X 省」→ 表内不得出现 X 以外的省。
      这正是漏 `--title` 时全国表被打成「广东省」的那道网。
    判据二（2026-10-13）：标题写「全国」→ 必须真覆盖 ≥ `NATIONAL_MIN_PROVINCES` 个省级行政区。
      事故原型：全库只有 **4 省 15 城**，成品却叫「全国政策补贴总表」——
      「数字会腐」那条判据管不到它（名字里没写计数），而**地理主张同样得有据**：
      报表使用者是机构（拿去报价、做方案），「全国」二字直接造成覆盖度误读。
    返回错误文案列表（空 = 通过）。
    """
    title = title or ""
    actual = sorted({M.province(r.get("city", "")) for r in recs} - {""})
    if "全国" in title and len(actual) < NATIONAL_MIN_PROVINCES:
        return ["标题自称「%s」，但实际只覆盖 %d 个省级行政区（%s）—— 未达 %d 个，"
                "「全国」与实际范围不符；请改用不含地理主张的名字（如「已收录城市」）"
                % (title, len(actual), "、".join(actual) or "无", NATIONAL_MIN_PROVINCES)]
    m = re.fullmatch(r"(.+?)省", title)
    if not m:
        return []
    declared = m.group(1)
    bad = [p for p in actual if p != declared]
    if bad:
        return [f"标题自称「{title}」，但表内含 {'、'.join(bad)} 的城市 —— 标题与实际范围不符"]
    return []


def build_rows(records):
    out = []
    for i, r in enumerate(records, 1):
        lo, hi = M.amount_low_high(r)
        out.append([
            i,
            r.get("id", ""),
            r.get("_scope_province") or M.province(r.get("city", "")),
            r.get("_scope_city") or M.city(r.get("city", "")),
            M.t("domain", r.get("domain", "")),
            r.get("sub_domain", "") or "",
            r.get("name", "") or "",
            r.get("std_name", "") or "",
            M.applicant_kinds(r.get("cond")),
            M.t("result_type", r.get("result_type", "")),
            M.t("approval_type", r.get("approval_type", "")),
            M.t("channel", r.get("channel", "")),
            M.amount_text(r),
            M.district_awards_text(r),
            M.apply_window_text(r),
            M.conditions_text(r),
            M.materials_text(r),
            M.entry_url(r),
            M.channel_note(r),
            M.source_doc(r),
            M.original_doc_url(r, AUDIT),
            M.deadline_text(r),
            M.t("depth", r.get("depth", "")),
            r.get("updated_at", "") or "",
        ])
    return out


# ── C 端行构造 / 排序 / 速览（2026-10-12）────────────────────────────────
def build_c_rows(records):
    """C 端行：与 `C_COLS` 一一对应。**空值一律空串**（`mappers._s()` 已单点收口 None）。"""
    out = []
    for i, r in enumerate(records, 1):
        out.append([
            i,
            r.get("id", ""),
            r.get("_scope_province") or M.province(r.get("city", "")),
            r.get("_scope_city") or M.city(r.get("city", "")),
            r.get("_scope_label") or "",
            M.t("domain", r.get("domain", "")),
            r.get("name", "") or "",
            M.applicant_kinds(r.get("cond")),
            M.t("result_type", r.get("result_type", "")),
            M.amount_text(r),
            M.conditions_text(r),
            M.materials_text(r),
            M.t("channel", r.get("channel", "")),
            M.entry_url(r),
            M.original_doc_url(r, AUDIT),
            M.source_doc(r),
            M.apply_window_text(r),
            M.deadline_text(r),
            M.t("depth", r.get("depth", "")),
            r.get("updated_at", "") or "",
        ])
    return out


def sort_c_rows(records, cities, multi_city):
    """C 端排序：**城市优先** —— 城市（条数降序 → 声明序 → 城市码）→ 类别 → 金额降序 → 名称。

    与 C 端 PDF 的「城市优先」同键，保证两侧条目顺序一致（客户拿 PDF 与 Excel 交叉核对不错位）。
    ★ 城市顺序**写死次序**：不依赖 Counter 的隐式顺序（字典序会给读者随机感）。
    ★ 本函数只排序、不筛条 —— 条数由 `recalc_rows()` 在成品矩阵侧现算。
    """
    if not multi_city:
        _dno = {d: i for i, d in enumerate(_domain_order())}
        lo_hi = {}
        for r in records:
            lo_hi[r.get("id")] = M.amount_low_high(r)
        return sorted(records, key=lambda r: (
            _dno.get(r.get("domain"), 99),
            (0 if lo_hi[r.get("id")][1] else 1, -(lo_hi[r.get("id")][1] or 0),
             str(r.get("name") or ""))))

    seq = list(cities or [])
    n_by_city = collections.Counter(r.get("city") for r in records)
    order = sorted(n_by_city, key=lambda c: (-n_by_city[c],
                                            seq.index(c) if c in seq else 99, str(c)))
    rank = {c: i for i, c in enumerate(order)}
    _dno = {d: i for i, d in enumerate(_domain_order())}
    def _key(r):
        lo, hi = M.amount_low_high(r)
        return (rank.get(r.get("city"), 99), _dno.get(r.get("domain"), 99),
                0 if hi else 1, -(hi or 0), str(r.get("name") or ""))
    return sorted(records, key=_key)


def _domain_order():
    """类别次序真值源 = `export_pdf.DOMAIN_ORDER`（不在此另抄一份）。

    函数内 import：`export_pdf` 顶层有 I/O 副作用（读 link_audit / 重设 stdout 编码），
    模块级 import 会污染 pytest 收集期（技能 §5i「模块顶层禁 I/O」）。
    """
    try:
        import export_pdf as EP
        return EP.DOMAIN_ORDER
    except Exception:
        return []


def build_city_overview(rows):
    """「按城市速览」行：城市 / 条数 / 类别分布 / 编号区间。

    这是多城清单的**第一跳索引** —— 读者先在这里看"有哪些城市、各多少条"，
    再回 Sheet1 用城市列筛选。单城产品只有一行，同样出（结构统一）。
    """
    order, idx, dom = [], collections.defaultdict(list), collections.defaultdict(collections.Counter)
    for i, r in enumerate(rows, 1):
        c = r.get("city") or ""
        if c not in idx:
            order.append(c)
        idx[c].append(i)
        dom[c][M.t("domain", r.get("domain", "")) or ""] += 1
    out = []
    for c in order:
        secs = "　".join("%s %d" % (k, v) for k, v in dom[c].most_common() if k)
        out.append([M.city(c) or c, len(idx[c]), secs,
                    "%d–%d" % (idx[c][0], idx[c][-1])])
    return out


def _width(s, cap):
    w = 0
    for ch in str(s):
        w += 2 if ord(ch) > 127 else 1
    return min(max(8, w + 2), cap)


def write_sheet(ws, rows, cols=None, center=None, freeze="C2", first_row=1):
    """写 Sheet1「政策总表」。

    `cols` / `center` / `freeze` 由**受众**决定（B 端 24 列、冻结「序号+政策编号」= C2；
    C 端 20 列、**冻结到「城市」列之后** = E2 —— 横向滚动看条件/材料时，
    城市列始终留在视野里，可以边筛城市边横向读内容。这正是弃用多城 PDF 的替代方案）。
    `first_row`：数据起始行（1 = 无前置说明行，表头在第 1 行）。
    """
    cols = COLS if cols is None else cols
    center = CENTER_COLS if center is None else center
    hr = first_row                       # 表头行号
    for c, (name, _w, _wrap) in enumerate(cols, 1):
        cell = ws.cell(row=hr, column=c, value=name)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[hr].height = 26

    widths = [len(n) + 2 for n, _, _ in cols]
    for ri, row in enumerate(rows, hr + 1):
        for ci, val in enumerate(row, 1):
            name = cols[ci - 1][0]
            cell = ws.cell(row=ri, column=ci, value=val)
            cell.border = BORDER
            wrap = cols[ci - 1][2]
            is_link = name in URL_COLS and str(val or "").startswith("http")
            if is_link:
                # L5：出成品前链接降级 —— 坏链/错链一律不写超链接，只留可读提示
                _role = URL_ROLE.get(name)
                _reason = LG.should_degrade(val, _role, AUDIT, LINK_STRICT) if _role else None
                if _reason:
                    DOWNGRADED.append((row[1] if len(row) > 1 else "", name, str(val), _reason))
                    # 客户可见文案必须中性，原因只进台账（不得暴露内部判据词）
                    cell.value = LG.CLIENT_TEXT.get(_role, "（链接暂不可用）")
                    val, is_link = cell.value, False
            cell.font = LINK_FONT if is_link else BODY_FONT
            if is_link:
                cell.hyperlink = str(val)          # 可点击打开原文
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=wrap,
                horizontal="center" if name in center else "left",
            )
            if (ri - hr) % 2 == 0:
                cell.fill = ALT_FILL
            if not wrap:
                widths[ci - 1] = max(widths[ci - 1], _width(val, cols[ci - 1][1]))
    for ci, (name, cap, _wrap) in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(ci)].width = min(max(widths[ci - 1], 8), cap)

    ws.freeze_panes = freeze
    ws.auto_filter.ref = f"A{hr}:{get_column_letter(len(cols))}{len(rows)+hr}"


def _parse_args(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--audience", default="b", choices=["b", "c"],
                    help="b=机构/决策者全量表（默认，24 列）；"
                         "c=**个人读者清单**（20 列）：只收个人/创业者/海外人才可申报、"
                         "城市列前置 + 「按城市速览」Sheet（见模块 docstring 的取舍说明）")
    ap.add_argument("--preset", default="",
                    help="b 端：gd / yrd / gx / all（省略=gd）；c 端：sz / prd / gd_prov / all")
    ap.add_argument("--cities", default="")
    ap.add_argument("--title", default="")
    ap.add_argument("--out", default="")
    ap.add_argument("--src", default=loader.DEFAULT_SRC)
    ap.add_argument("--scope", default="region", choices=["region", "province"],
                    help="（仅 b 端）region=区域/城市包（默认，原文照出）；"
                         "province=**省总表**：启用跨城归并（省级/国家级政策合并为一条并标注适用范围）"
                         "+ 覆盖率标注，要求全部城市同属一省")
    ap.add_argument("--link-strict", action="store_true",
                    help="L5：连栏目/聚合页、政策解读页、新闻页也降级（默认只降死链+明文文种错）")
    return ap.parse_args(argv)


def main():
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

    args = _parse_args()

    global LINK_STRICT
    LINK_STRICT = args.link_strict

    if args.audience == "c":
        return main_c(args)
    return main_b(args)


def main_b(args):
    """B 端：全量工具表（机构 / 决策者）。"""
    if args.preset and args.preset not in loader.PRESETS:
        print(f"❌ b 端 --preset 必须是 {' / '.join(loader.PRESETS)} 之一，收到 {args.preset!r}")
        sys.exit(2)
    args = argparse.Namespace(**dict(vars(args), preset=(args.preset or "gd")))
    cities = [c.strip() for c in args.cities.split(",") if c.strip()] or loader.PRESETS.get(args.preset)
    all_rows = loader.load_all(args.src)
    recs = loader.filter_cities(all_rows, cities)

    # ── 省总表模式（--scope province）────────────────────────────────
    # 与「区域包」并存而不替换：区域包按经济圈（长三角跨省）、省总表按行政边界。
    _prov, _merge_stat, _cov = None, None, None
    if args.scope == "province":
        _provs = sorted({M.province(c) for c in (cities or [])} - {""})
        if len(_provs) != 1:
            print(f"❌ --scope province 要求全部城市同属一省，实际 {_provs or '（未指定城市）'}"
                  f"；跨省请用 --preset yrd 之类的区域包。")
            sys.exit(2)
        _prov = _provs[0]
        # 标题交给 derive_title(scope="province") 统一推导，此处不重复设值

    # 成品层收录口径（2026-10-10）：**只接入现行有效政策**。
    # 非现役（exited/renamed/merged/superseded）仍留在库、前端仍以「备查」区块展示，
    # 但不进售卖成品 —— 成品是"你能申报的政策"清单，放进去只会制造无效申报。
    _n_city = len(recs)
    recs = loader.filter_active(recs)
    _n_dropped = _n_city - len(recs)

    # 跨城归并（省总表专属）：省级/国家级政策被各市各录一条 → 合并为一条并标注适用范围。
    # 判据真值源 = pipeline/cross_city.py（区域包不做归并，原文照出）。
    if _prov:
        # 传**全库**判层级：省表只含广东城市，本表内看不出「跨省 → 国家级」
        recs, _merge_stat = cross_city.merge(recs, province=_prov,
                                             all_rows=loader.filter_active(all_rows))
        _cov = geo.coverage(_prov, cities)

    # 排序：省份 -> 城市 -> 类别 -> 编号，让同城同类的政策排一起
    recs.sort(key=lambda r: (M.province(r.get("city", "")), M.city(r.get("city", "")),
                             M.t("domain", r.get("domain", "")), str(r.get("id", ""))))

    if not recs:
        # ★ 0 条时**不出空表**：只有表头的 Excel 会被当成成品发出去，
        #   客户点开一片空白，比直接报错更坏（也是「不静默降级」的同一条纪律）。
        print("❌ 这个范围里没有任何现行有效政策，不出空表。")
        print("   · 检查 --src / --cities / --preset 是否指对了城市（当前范围：%s）"
              % (",".join(cities) if cities else "全库"))
        print("   · 库里确实没有这座城的数据：先按 policy_hub/domains/README.md 采集填入")
        print("   · 有数据但全被过滤：这些条目的 valid 可能都不是 active")
        sys.exit(2)

    title = derive_title(args.title, args.preset, cities, recs, args.scope)
    _scope_err = check_title_scope(title, recs)
    if _scope_err:
        for e in _scope_err:
            print(f"❌ {e}")
        print("   （这是 2026-10-10 修掉的「优先级陷阱」重现 —— 标题必须按声明范围推导）")
        sys.exit(2)
    out = args.out or os.path.join(
        os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "output",
        f"{title}政策补贴总表_{VERSION}.xlsx")
    os.makedirs(os.path.dirname(out), exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "政策总表"
    rows = build_rows(recs)
    write_sheet(ws, rows)

    # Sheet2 字段说明
    ws2 = wb.create_sheet("字段说明")
    ws2.append(["列名", "说明"])
    for c in ws2[1]:
        c.fill, c.font = HEAD_FILL, HEAD_FONT
    for n, note in FIELD_NOTE:
        ws2.append([n, note])
    ws2.column_dimensions["A"].width = 16
    ws2.column_dimensions["B"].width = 70
    for row in ws2.iter_rows(min_row=2):
        row[0].font = Font(bold=True, size=10, name="微软雅黑")
        row[1].font = BODY_FONT
        row[1].alignment = Alignment(wrap_text=True, vertical="top")

    # Sheet3 数据字典 + 免责声明
    ws3 = wb.create_sheet("数据字典与声明")
    ws3.append(["字段", "原始值", "中文含义"])
    for c in ws3[1]:
        c.fill, c.font = HEAD_FILL, HEAD_FONT
    for r in DICT_ROWS:
        ws3.append(list(r))
    ws3.append([])
    _orig = [M.original_doc_url(r, AUDIT) for r in recs]
    _ent = [M.entry_url(r) for r in recs]
    src_uniq = sum(1 for u, e in zip(_orig, _ent) if u and u != e)
    src_same = sum(1 for u, e in zip(_orig, _ent) if u and u == e)
    src_none = sum(1 for u in _orig if not u)
    notes = [
        ["表结构版本", SHEET_VERSION],
        ["数据版本", VERSION],
        ["数据来源", "各地人力资源和社会保障局、税务局等政府门户网站公开信息（官方域名）"],
        ["统计口径", f"共 {len(recs)} 条政策，覆盖 {len(set(r.get('city') for r in recs))} 个城市"
                    f"（仅收录现行有效政策：已废止 / 已更名 / 已合并 / 已被新版替代的政策"
                    f"本表不收录，如需备查请使用在线版本）"],
    ]
    # ── 省总表专属：覆盖率标注 + 跨城归并说明（2026-10-10 用户拍板）──────
    # 覆盖率必须诚实：只采了省内部份城市时，出「某省总表」若不说清楚，
    # 客户会以为覆盖全省 —— 这是**减少售后纠纷**的关键一行，不是装饰。
    if _cov:
        _cv = "、".join(_cov["covered_names"])
        if _cov["total"]:
            notes.append(["覆盖范围",
                          f"{_prov}共 {_cov['total']} 个地级市，本表收录 {len(_cov['covered'])} 个：{_cv}"])
            if _cov["missing_names"]:
                notes.append(["未覆盖城市",
                              f"{'、'.join(_cov['missing_names'])}（共 {len(_cov['missing_names'])} 个，"
                              f"尚未采集，本表不含其政策）"])
        else:
            notes.append(["覆盖范围", f"本表收录 {len(_cov['covered'])} 个城市：{_cv}"])
    if _merge_stat:
        notes.append(["跨城归并",
                      f"省级与国家级政策被多个城市重复收录的，已合并为一条并在「城市」列标注适用范围"
                      f"（省级＝「{_prov}全省」、国家级＝「全国」）；本次合并 {_merge_stat['groups']} 组、"
                      f"去重 {_merge_stat['dropped']} 条重复记录。"])
    notes += [
        ["溯源覆盖", f"{src_uniq}/{len(recs)} 条含可点击的政策原文链接（独立于办理入口）；"
                    f"{src_same} 条原文与办事入口为同一页面；"
                    f"{src_none} 条暂未检索到公开发布的原文页（仅标注文号）"],
        ["生成日期", datetime.date.today().isoformat()],
        ["免责声明",
         "本表基于政府部门公开发布的政策文件整理，标注「规则估算」的金额为按文件口径推算，"
         "实际补贴标准、申报条件与窗口期一律以官方最新文件及经办窗口审核结果为准。"
         "本产品提供信息整理服务，不代办申报、不承诺审批结果。"],
    ]
    for k, v in notes:
        ws3.append([k, v])
        ws3.cell(row=ws3.max_row, column=1).font = Font(bold=True, size=10, name="微软雅黑")
        ws3.cell(row=ws3.max_row, column=2).font = BODY_FONT
        ws3.cell(row=ws3.max_row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    ws3.column_dimensions["A"].width = 16
    ws3.column_dimensions["B"].width = 22
    ws3.column_dimensions["C"].width = 60

    wb.save(out)

    if DOWNGRADED:
        _dp = os.path.join(os.path.dirname(out),
                           "_link_degrade_%s.md" % datetime.datetime.now().strftime("%Y%m%d-%H%M%S"))
        _lines = ["# 成品链接降级台账（L5）", "",
                  "> 成品：`%s`" % os.path.basename(out),
                  "> 规则：`pipeline/link_guard.py`｜判据源：源库 `policy_hub/collector/link_audit.json`",
                  "> 含义：下列链接**未写入成品**（客户看不到超链接），只留可读提示，避免点到坏链/错页。", ""]
        _lines += ["- `%s` ｜ %s ｜ %s ｜ %s" % (r, c, why, u) for r, c, u, why in DOWNGRADED]
        open(_dp, "w", encoding="utf-8").write("\n".join(_lines))
        print("🔗 链接降级 %d 处（客户不会点到坏链）→ %s" % (len(DOWNGRADED), _dp))

    print(f"✅ 已导出：{out}")
    print(f"   记录数：{len(recs)}  |  城市：{len(set(r.get('city') for r in recs))}  |  列数：{len(COLS)}")
    if _n_dropped:
        print(f"   ℹ️ 已过滤非现行有效政策 {_n_dropped} 条（成品只收录现行有效；"
              f"需备查请用在线版本，那里保留「已退出实施·备查」区块）")
    if _merge_stat:
        print(f"   🔀 跨城归并 {_merge_stat['groups']} 组、去重 {_merge_stat['dropped']} 条"
              f"（省级→「{_prov}全省」、国家级→「全国」）")
        if _merge_stat["divergent"]:
            print(f"   ⚠️ {len(_merge_stat['divergent'])} 组金额表述不一致 —— 已保留代表记录，"
                  f"明细见台账（供人工复核，未静默取舍）")
        _mp = os.path.join(os.path.dirname(out),
                           "_cross_city_%s.md" % datetime.datetime.now().strftime("%Y%m%d-%H%M%S"))
        _ml = ["# 跨城归并台账（省总表）", "",
               "> 成品：`%s`" % os.path.basename(out),
               "> 判据：`pipeline/cross_city.py`｜键 = 归一化（政策名称 + 政策依据文号）",
               "> 含义：下列为**被合并掉**的重复记录组，成品中只保留 1 条并标注适用范围。", ""]
        for d in sorted(_merge_stat["detail"], key=lambda x: -x["n"]):
            _ml.append("- `%s` ｜ %s ｜ %d 城（%s）｜ 保留 `%s`"
                       % (d["key"], d["scope"], len(d["cities"]), "、".join(d["cities"]), d["kept"]))
        if _merge_stat["divergent"]:
            _ml += ["", "## 金额表述不一致（未合并取舍，保留代表记录，请人工复核）", ""]
            for d in _merge_stat["divergent"]:
                _ml.append("- `%s` ｜ %s ｜ %d 种表述" % (d["key"], d["scope"], d["n_amount_variants"]))
        open(_mp, "w", encoding="utf-8").write("\n".join(_ml))
        print(f"   📋 归并明细 → {_mp}")
    if _cov and _cov["total"]:
        print(f"   📍 覆盖 {_prov} {len(_cov['covered'])}/{_cov['total']} 个地级市，"
              f"未覆盖 {len(_cov['missing_names'])} 个（已在表内标注）")
    print()
    if loader.MISMATCH:
        print(f"⚠️ 源数据字段与目录不一致（本表已按目录归位）：{len(loader.MISMATCH)} 处")
        for rid, fld, got, want, src in loader.MISMATCH[:12]:
            print(f"   - {rid}: {fld}={got!r} → 应为 {want!r}   [{src}]")
        if len(loader.MISMATCH) > 12:
            print(f"   … 另有 {len(loader.MISMATCH)-12} 处")
        print()
    print(M.report_unmapped())


def main_c(args):
    """C 端：**个人可申报清单（Excel）**。

    为什么多城是 Excel（2026-10-12 用户实测反馈）：
      单城 63 条用 PDF 刚好；珠三角 255 条 / 全国 385 条时，PDF 要线性翻完才知道有没有
      自己要的政策，且"先记政策名、再找城市"两维人工扫 → 改 Excel 用城市列筛选定位。

    产品定义（标题 / 城市集 / 文件名）真值源 = `export_pdf.PRODUCTS` —— 与 C 端 PDF 同一份，
    不在这里另建产品表（否则"全国到底几个城市"会出现两个答案）。
    收录口径真值源 = `audience.is_person`（同上）。
    """
    import export_pdf as EP                       # 函数内 import：其顶层有 I/O（技能 §5i）

    try:
        # ★ 支持 `city_<城市码>`（单城上架单元）—— 单城产品由 `EP.city_product()` 现算，
        #   不在静态 `PRODUCTS` 里，所以必须走 `resolve_preset` 而不是直接下标（2026-10-14）。
        prod = EP.resolve_preset(args.preset)
    except KeyError:
        print("❌ --audience c 的 --preset 必须是 %s 之一，或 city_<城市码>（如 city_sz），收到 %r"
              % (" / ".join(EP.PRODUCTS), args.preset))
        sys.exit(2)

    all_rows = loader.load_all(args.src)
    rows = loader.filter_cities(all_rows, prod["cities"])
    _n_city = len(rows)
    # 成品层收录口径：只接现行有效（与 B 端/C 端 PDF 同源判据）
    rows = loader.filter_active(rows)
    _n_dropped = _n_city - len(rows)
    _active_all = loader.filter_active(all_rows)

    _n_active = len(rows)
    rows = [r for r in rows if is_person(r)]
    _n_imperson = _n_active - len(rows)

    if not rows:
        print("❌ 这个范围里没有可给个人/创业者/海外人才申报的现行政策，不出空表。")
        print("   · 库里可能只有单位类政策 —— C 端清单只收个人能申报的")
        print("   · 或者这座城市还没填数据：见 policy_hub/domains/README.md")
        sys.exit(2)

    cities = prod.get("cities")
    multi_city = (cities is None) or (len(cities) > 1)
    if multi_city:
        # 「适用范围」文案真值源 = cross_city.scope_labels（与 C 端 PDF 同源）
        _labels = cross_city.scope_labels(rows, _active_all, prod.get("province"))
        for r in rows:
            r["_scope_label"] = _labels.get(r.get("id")) or ""

    rows = sort_c_rows(rows, cities, multi_city)

    title = prod["fname"]
    out = args.out or os.path.join(
        os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "output", "C端",
        f"{title}_{VERSION}.xlsx")
    os.makedirs(os.path.dirname(out), exist_ok=True)

    print(f"=== {prod['title']}（C 端 · Excel）===")
    print(f"源读取：{_n_city} 条 → 现行有效 {_n_active} 条（过滤非现役 {_n_dropped} 条）"
          f" → 个人可申报 {len(rows)} 条（滤掉非个人条目 {_n_imperson} 条）")
    print(f"覆盖：{len(set(r.get('city') for r in rows))} 个城市 ｜ 列数 {len(C_COLS)}"
          f" ｜ 冻结窗格 E2（横向滚动时「城市」列常驻可见）")

    wb = Workbook()
    ws = wb.active
    ws.title = "政策总表"
    write_sheet(ws, build_c_rows(rows), cols=C_COLS, center=C_CENTER_COLS, freeze="E2")

    # ── Sheet2 按城市速览（多城清单的第一跳索引）────────────────────────
    ov = build_city_overview(rows)
    ws2 = wb.create_sheet("按城市速览")
    ws2.append(["城市", "政策条数", "补贴类别分布", "编号区间"])
    for c in ws2[1]:
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        c.alignment = Alignment(horizontal="center", vertical="center")
    for r in ov:
        ws2.append(r)
        ws2.cell(row=ws2.max_row, column=2).alignment = Alignment(horizontal="center")
        ws2.cell(row=ws2.max_row, column=4).alignment = Alignment(horizontal="center")
        ws2.cell(row=ws2.max_row, column=3).font = BODY_FONT
    ws2.append([])
    ws2.append(["合计", len(rows), f"{len(ov)} 个城市", ""])
    ws2.cell(row=ws2.max_row, column=1).font = Font(bold=True, size=10, name="微软雅黑")
    ws2.append([])
    # 「怎么用」文案按**单城 / 多城**分叉：单城清单点「城市」筛选毫无意义（整本同城）。
    # ★ 并明确写出「两列链接可点、可复制」—— 这正是使用者最常问的。
    _howto = ("切到「政策总表」：按「政策类别」筛大类，用「数据深度」筛 L3/L4 只看信息最完整的条目。"
              "★「办理入口」「政策原文链接」两列是**可直接点击**的链接，也可**复制**到浏览器打开。")
    if multi_city:
        _howto = "切到「政策总表」，点「城市」列的筛选箭头勾选城市 → 只看该城政策；" + _howto
    ws2.append(["怎么用", _howto, "", ""])
    ws2.cell(row=ws2.max_row, column=1).font = Font(bold=True, size=10, name="微软雅黑")
    ws2.cell(row=ws2.max_row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    ws2.column_dimensions["A"].width = 14
    ws2.column_dimensions["B"].width = 12
    ws2.column_dimensions["C"].width = 52
    ws2.column_dimensions["D"].width = 12

    # ── Sheet3 字段说明 ────────────────────────────────────────────────
    ws3 = wb.create_sheet("字段说明")
    ws3.append(["列名", "说明"])
    for c in ws3[1]:
        c.fill, c.font = HEAD_FILL, HEAD_FONT
    for n, note in C_FIELD_NOTE:
        ws3.append([n, note])
    ws3.column_dimensions["A"].width = 16
    ws3.column_dimensions["B"].width = 76
    for row in ws3.iter_rows(min_row=2):
        row[0].font = Font(bold=True, size=10, name="微软雅黑")
        row[1].font = BODY_FONT
        row[1].alignment = Alignment(wrap_text=True, vertical="top")

    # ── Sheet4 数据字典与声明 ──────────────────────────────────────────
    ws4 = wb.create_sheet("数据字典与声明")
    ws4.append(["字段", "原始值", "中文含义"])
    for c in ws4[1]:
        c.fill, c.font = HEAD_FILL, HEAD_FONT
    for r in DICT_ROWS:
        ws4.append(list(r))
    ws4.append([])
    _orig = [M.original_doc_url(r, AUDIT) for r in rows]
    _ent = [M.entry_url(r) for r in rows]
    src_uniq = sum(1 for u, e in zip(_orig, _ent) if u and u != e)
    src_same = sum(1 for u, e in zip(_orig, _ent) if u and u == e)
    src_none = sum(1 for u in _orig if not u)
    _n_city_cov = len(set(r.get("city") for r in rows))
    # ★ 单城 / 多城的口径必须分叉 —— 单城清单讲「各市分别列出」是错的，
    #   而「适用范围」列在单城下**整列留空**（不解释一句，客户会以为是没做完）。
    _scope_dd = (
        "多城清单里，省级与国家级政策各市会各自发布一次 → 本表**按城市分别列出、未做归并**，"
        "靠「适用范围」列区分：标「某省全省」「全国通用」的是上级政策，"
        "在多个城市下都会出现，**不是重复条目**。") if multi_city else (
        "本表为**单城清单**，整本都是同一座城市，故「适用范围」列**整列留空**"
        "（逐条重复标城市名是噪音）。")
    _howto_dd = (
        "① 在「按城市速览」看有哪些城市、各多少条；② 回「政策总表」点「城市」列筛选箭头选城市；"
        "③ 需要更完整的条款时，点「政策原文链接」或「办理入口」打开官方页面核对。") if multi_city else (
        "① 在「按城市速览」看本城政策条数与类别分布；② 回「政策总表」按「政策类别」筛大类；"
        "③ 需要更完整的条款时，点「政策原文链接」或「办理入口」打开官方页面核对。")
    notes = [
        ["表结构版本", SHEET_VERSION],
        ["数据版本", VERSION],
        ["数据来源", "各地人力资源和社会保障局、税务局、住房和城乡建设等主管部门门户网站公开信息（官方域名）"],
        ["收录口径", f"共 {len(rows)} 条，覆盖 {_n_city_cov} 个城市；**只收录「个人 / 创业者 / "
                    f"{_OS_KIND}」可申报的政策**（企业专项不进本表）。仅收录现行有效政策：已废止 / "
                    f"已更名 / 已合并 / 已被新版替代的政策本表不含，如需备查请使用在线版本。"],
        ["适用范围", _scope_dd],
        ["怎么用", _howto_dd],
        ["链接可用性", "「办理入口」与「政策原文链接」两列：**由官方域名构成，单元格本身即可点击**；"
                      "若你的阅读器/软件不支持点击，**复制该单元格内容**粘到浏览器同样能打开。"],
        ["溯源覆盖", f"{src_uniq}/{len(rows)} 条含可点击的政策原文链接（独立于办理入口）；"
                    f"{src_same} 条原文与办事入口为同一页面；"
                    f"{src_none} 条暂未检索到公开发布的原文页（仅标注文号）"],
        ["生成日期", datetime.date.today().isoformat()],
        ["免责声明",
         "本表基于政府部门公开发布的政策文件整理，标注「规则估算」的金额为按文件口径推算，"
         "实际补贴标准、申报条件与窗口期一律以官方最新文件及经办窗口审核结果为准。"
         "本产品提供信息整理服务，不代办申报、不承诺审批结果。"],
    ]
    for k, v in notes:
        ws4.append([k, v])
        ws4.cell(row=ws4.max_row, column=1).font = Font(bold=True, size=10, name="微软雅黑")
        ws4.cell(row=ws4.max_row, column=2).font = BODY_FONT
        ws4.cell(row=ws4.max_row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    ws4.column_dimensions["A"].width = 16
    ws4.column_dimensions["B"].width = 24
    ws4.column_dimensions["C"].width = 60

    wb.save(out)

    if DOWNGRADED:
        _dp = os.path.join(os.path.dirname(out),
                           "_link_degrade_%s.md" % datetime.datetime.now().strftime("%Y%m%d-%H%M%S"))
        _lines = ["# 成品链接降级台账（L5）", "",
                  "> 成品：`%s`" % os.path.basename(out),
                  "> 规则：`pipeline/link_guard.py`｜判据源：源库 `policy_hub/collector/link_audit.json`",
                  "> 含义：下列链接**未写入成品**（客户看不到超链接），只留可读提示，避免点到坏链/错页。", ""]
        _lines += ["- `%s` ｜ %s ｜ %s ｜ %s" % (r, c, why, u) for r, c, u, why in DOWNGRADED]
        open(_dp, "w", encoding="utf-8").write("\n".join(_lines))
        print("🔗 链接降级 %d 处（客户不会点到坏链）→ %s" % (len(DOWNGRADED), _dp))

    print(f"✅ 已导出：{out}")
    print(f"   记录数：{len(rows)}  |  城市：{len(set(r.get('city') for r in rows))}"
          f"  |  列数：{len(C_COLS)}  |  Sheet：政策总表 / 按城市速览 / 字段说明 / 数据字典与声明")
    if loader.MISMATCH:
        print(f"⚠️ 源数据字段与目录不一致（本表已按目录归位）：{len(loader.MISMATCH)} 处")
    print()
    print(M.report_unmapped())


if __name__ == "__main__":
    main()
